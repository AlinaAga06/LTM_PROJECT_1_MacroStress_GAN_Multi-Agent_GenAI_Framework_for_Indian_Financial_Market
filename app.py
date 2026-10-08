"""
MacroStress-GAN | Financial Risk Dashboard  (dynamic edition)

Same dashboard layout as before, but the uploaded document is the ONLY source of truth.
Nothing is read from old scenario CSVs, old optimizer JSONs or hard-coded weights/sensitivities/shocks.

Workflow: Document -> Parse & validate -> TimeGAN scenarios (1,000 x N days)
          -> Day-by-day portfolio risk -> VaR / ES / Expected Loss / Drawdown
          -> Portfolio optimization -> Re-stress verification -> Final report

Run:  streamlit run app.py
"""
from __future__ import annotations

import importlib
import io
import json
import re
import sys
import traceback
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import plotly.graph_objects as go
import streamlit as st

PROJECT_ROOT = Path(__file__).resolve().parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

APP_NAME, APP_VERSION = "MacroStress-GAN Risk Platform", "4.3"

# ============================================================
# CONSTANTS
# ============================================================
STOCK_FEATURES = ["NIFTY50_Return", "CRUDE_OIL_Return", "USD_INR_Return",
                  "INDIA_VIX_Change", "INDIA_10Y_YIELD_Change"]
RISK_VARIABLES = ["NIFTY50", "CRUDE_OIL", "USD_INR", "INDIA_VIX", "INDIA_10Y_YIELD"]
RISK_TO_INDEX = {v: i for i, v in enumerate(RISK_VARIABLES)}
RISK_KEYS = ["VaR", "ES", "Expected Loss", "Max Drawdown", "Worst Loss"]
NUM_SCENARIOS, DEFAULT_SEED = 1000, 42

# ---- Colour system -------------------------------------------------------
NAVY, BLUE, TEAL, GREEN = "#0F2A47", "#2563EB", "#0D9488", "#16A34A"
AMBER, ORANGE, RED, SLATE, MUTED = "#F59E0B", "#EA580C", "#DC2626", "#475569", "#94A3B8"
ACCENTS = {
    "blue": (BLUE, "#EFF6FF"), "teal": (TEAL, "#F0FDFA"), "green": (GREEN, "#F0FDF4"),
    "amber": (AMBER, "#FFFBEB"), "orange": (ORANGE, "#FFF7ED"), "red": (RED, "#FEF2F2"),
    "slate": (SLATE, "#F8FAFC"),
}
RATING_ACCENT = {"Low": "green", "Moderate": "amber", "High": "orange", "Severe": "red"}

st.set_page_config(page_title=APP_NAME, page_icon="📊", layout="wide", initial_sidebar_state="expanded")

SESSION_DEFAULTS = {"config": None, "upload_key": None, "upload_name": None, "simulation": None,
                    "optimization": None, "restress": None, "pipeline_error": None}
for _k, _v in SESSION_DEFAULTS.items():
    st.session_state.setdefault(_k, _v)

CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');
html, body, [class*="css"] {font-family:"Inter","Segoe UI",Arial,sans-serif;}
.stApp {background:#F1F5F9;}
.block-container {padding-top:1.4rem; padding-bottom:2.5rem; max-width:1500px;}

/* Sidebar: dark navy, light labels, dark text inside inputs */
section[data-testid="stSidebar"] {background:linear-gradient(180deg,#0F2A47 0%,#143D66 100%);}
section[data-testid="stSidebar"] h1, section[data-testid="stSidebar"] h2,
section[data-testid="stSidebar"] h3, section[data-testid="stSidebar"] hr {color:#FFFFFF !important;}
section[data-testid="stSidebar"] [data-testid="stWidgetLabel"] p,
section[data-testid="stSidebar"] [data-testid="stMarkdownContainer"] p,
section[data-testid="stSidebar"] [data-testid="stCaptionContainer"] {color:#CBD5E1 !important;}
section[data-testid="stSidebar"] input {color:#0F172A !important;}

/* Hero */
.hero-box {padding:2rem 2.2rem 1.7rem 2.2rem; border-radius:18px; color:#fff; margin-bottom:1.3rem;
           background:linear-gradient(120deg,#0F2A47 0%,#1E4E8C 55%,#0D9488 100%);
           box-shadow:0 6px 18px rgba(15,42,71,.18);}
.hero-box h1 {color:#fff; font-size:2.3rem; margin:0 0 .4rem 0; font-weight:800;}
.hero-box p {color:#DBEAFE; font-size:1.02rem; margin:0;}

/* KPI cards */
.kpi-card {border-radius:14px; padding:1rem 1.1rem; min-height:118px;
           box-shadow:0 2px 8px rgba(15,23,42,.06); border:1px solid #E2E8F0;}
.kpi-label {color:#475569; font-size:.74rem; font-weight:700; letter-spacing:.07em; text-transform:uppercase;}
.kpi-value {color:#0F172A; font-size:1.6rem; font-weight:800; margin-top:.3rem; line-height:1.15; word-break:break-word;}
.kpi-sub {color:#475569; font-size:.8rem; font-weight:500; margin-top:.3rem;}

/* Info cards / steps */
.info-card {background:#fff; border:1px solid #E2E8F0; border-radius:14px; padding:1.1rem 1.25rem;
            box-shadow:0 2px 8px rgba(15,23,42,.05); height:100%;}
.info-card h3 {margin:0 0 .3rem 0; color:#0F172A; font-size:1.2rem;}
.info-card p {margin:.15rem 0; color:#475569; font-size:.9rem;}
.step {display:flex; align-items:center; gap:.7rem; background:#fff; border:1px solid #E2E8F0;
       border-radius:12px; padding:.7rem .9rem; margin-bottom:.6rem; font-weight:600; color:#1E293B; font-size:.92rem;}
.step-no {min-width:28px; height:28px; border-radius:50%; background:#2563EB; color:#fff;
          display:flex; align-items:center; justify-content:center; font-size:.82rem; font-weight:800;}
.check {padding:9px 14px; margin:5px 0; border-radius:8px; background:#fff;
        border:1px solid #E2E8F0; font-size:.9rem; font-weight:600; color:#334E68;}
.section-title {color:#0F2A47; font-size:1.4rem; font-weight:800; margin:.9rem 0 .2rem 0;
                padding-left:.7rem; border-left:5px solid #2563EB;}

/* Tabs */
button[data-baseweb="tab"] p {color:#475569 !important; font-weight:600; font-size:.98rem;}
button[data-baseweb="tab"][aria-selected="true"] p {color:#2563EB !important;}
div[data-baseweb="tab-highlight"] {background-color:#2563EB !important; height:3px;}
div[data-baseweb="tab-border"] {background-color:#CBD5E1 !important;}

/* Buttons */
.stDownloadButton button, .stButton button {background:#2563EB; color:#fff !important; border:0;
        border-radius:9px; font-weight:600; padding:.5rem 1.1rem;}
.stDownloadButton button:hover, .stButton button:hover {background:#1D4ED8; color:#fff !important;}

[data-testid="stDataFrame"] {border:1px solid #E2E8F0; border-radius:10px;}
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)


# ============================================================
# UI HELPERS
# ============================================================
def pct(value: Any, digits: int = 2) -> str:
    try:
        v = float(value)
        return f"{v * 100:.{digits}f}%" if np.isfinite(v) else "N/A"
    except Exception:
        return "N/A"


def money(value: Any, currency: str = "INR") -> str:
    try:
        v = float(value)
        if not np.isfinite(v):
            return "N/A"
        if currency.upper() == "INR":
            if abs(v) >= 1e7:
                return f"₹{v / 1e7:.2f} Cr"
            if abs(v) >= 1e5:
                return f"₹{v / 1e5:.2f} L"
            return f"₹{v:,.2f}"
        return f"{currency.upper()} {v:,.2f}"
    except Exception:
        return "N/A"


def label(key: str, conf: float) -> str:
    return {"VaR": f"VaR {conf:.0%}", "ES": f"Expected Shortfall {conf:.0%}"}.get(key, key)


def kpi(lbl: str, value: str, sub: str = "", accent: str = "blue") -> None:
    border, tint = ACCENTS.get(accent, ACCENTS["blue"])
    st.markdown(f'<div class="kpi-card" style="background:{tint}; border-top:4px solid {border};">'
                f'<div class="kpi-label">{lbl}</div><div class="kpi-value">{value}</div>'
                f'<div class="kpi-sub">{sub}</div></div>', unsafe_allow_html=True)


def title(text: str, subtitle: str | None = None) -> None:
    st.markdown(f'<div class="section-title">{text}</div>', unsafe_allow_html=True)
    if subtitle:
        st.caption(subtitle)


def risk_rating(value: float) -> str:
    return "Low" if value < 0.05 else "Moderate" if value < 0.10 else "High" if value < 0.20 else "Severe"


def style_fig(fig: go.Figure, height: int = 400) -> go.Figure:
    fig.update_layout(height=height, margin=dict(l=20, r=20, t=55, b=30),
                      paper_bgcolor="white", plot_bgcolor="white", hovermode="x unified",
                      font=dict(family="Inter, Arial", color="#0F172A"),
                      title_font=dict(size=16, color=NAVY),
                      colorway=[BLUE, TEAL, AMBER, RED, GREEN],
                      legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1))
    fig.update_xaxes(showgrid=True, gridcolor="#E2E8F0", zeroline=False)
    fig.update_yaxes(showgrid=True, gridcolor="#E2E8F0", zeroline=False)
    return fig


def yield_to_decimal(v: float) -> float:
    """Yield shocks follow the SimulationAgent convention: |v| >= 0.1 means percentage points."""
    return v / 100.0 if abs(v) >= 0.1 else v


def shock_display(variable: str, v: float | None) -> str:
    if v is None or not np.isfinite(v):
        return "Missing"
    if variable == "INDIA_10Y_YIELD":
        return f"{yield_to_decimal(v) * 100:+.2f} pp"
    return f"{v * 100:+.2f}%"


# ============================================================
# DOCUMENT READING
# ============================================================
def read_uploaded_file(uploaded_file) -> str:
    name, raw = uploaded_file.name.lower(), uploaded_file.getvalue()
    if name.endswith((".txt", ".py", ".md")):
        return raw.decode("utf-8", errors="ignore")
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(raw)).pages)
    if name.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(raw))
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(lines)
    raise ValueError("Unsupported file type. Upload PDF, DOCX, TXT, PY or MD.")


# ============================================================
# PARSING
# ============================================================
def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return re.sub(r"<[^>]+>", " ", text)


def extract_number(text: Any) -> float | None:
    m = re.search(r"[-+]?\d[\d,]*(?:\.\d+)?", str(text))
    try:
        return float(m.group(0).replace(",", "")) if m else None
    except Exception:
        return None


def parse_percent_or_decimal(s: Any) -> float | None:
    s = str(s).strip()
    n = extract_number(s)
    if n is None:
        return None
    if "%" in s:
        return n / 100.0
    if abs(n) <= 1.0:
        return n
    return n / 100.0 if abs(n) <= 100.0 else n


def parse_shock_value(s: Any, variable: str) -> float | None:
    """Market variables -> fraction (-15% -> -0.15). Yield -> raw number (agent normalizes)."""
    s = str(s).strip()
    n = extract_number(s)
    if n is None:
        return None
    if variable == "INDIA_10Y_YIELD":
        return n
    if "%" in s or abs(n) > 1:
        return n / 100.0
    return n


def find_first(patterns: list[str], text: str, flags=re.IGNORECASE | re.MULTILINE):
    for p in patterns:
        m = re.search(p, text, flags)
        if m:
            return m
    return None


VARIABLE_ALIASES = {
    "NIFTY50": [r"NIFTY\s*50", r"NIFTY50"],
    "CRUDE_OIL": [r"CRUDE[_\s\-]?OIL", r"CRUDEOIL"],
    "USD_INR": [r"USD[_\s/\-]?INR", r"USDINR"],
    "INDIA_VIX": [r"INDIA[_\s\-]?VIX", r"\bVIX\b"],
    "INDIA_10Y_YIELD": [r"INDIA[_\s\-]?10Y[_\s\-]?YIELD", r"10Y[_\s\-]?YIELD", r"INDIA[_\s\-]?10Y"],
}


def normalize_key(value: str) -> str:
    s = re.sub(r"_+", "_", re.sub(r"[/\-\s]", "_", str(value).upper().strip()))
    aliases = {"NIFTY": "NIFTY50", "NIFTY_50": "NIFTY50", "CRUDE": "CRUDE_OIL", "CRUDEOIL": "CRUDE_OIL",
               "USDINR": "USD_INR", "VIX": "INDIA_VIX", "INDIAVIX": "INDIA_VIX",
               "10Y_YIELD": "INDIA_10Y_YIELD", "INDIA10Y_YIELD": "INDIA_10Y_YIELD", "INDIA_10Y": "INDIA_10Y_YIELD"}
    return aliases.get(s, s)


def parse_institution_type(text: str) -> str:
    m = find_first([r"Institution\s*Type\s*[:=]\s*([^\n]+)", r"institution_type\s*[:=]\s*[\"']?([^\"'\n,}]+)",
                    r"Institution\s*[:=]\s*([^\n]+)"], text)
    return m.group(1).strip().strip("\"'") if m else "Stock Market"


def parse_currency(text: str) -> str:
    m = find_first([r"Currency\s*[:=]\s*([A-Za-z]{3})", r"Portfolio\s*Amount\s*[:=]\s*(INR|USD|EUR|GBP)"], text)
    return m.group(1).upper() if m else "INR"


def parse_portfolio_amount(text: str) -> float | None:
    m = find_first([
        r"Portfolio\s*Amount\s*[:=]\s*(?:INR|USD|EUR|GBP|₹|Rs\.?)?\s*([\d,]+(?:\.\d+)?)",
        r"portfolio_amount\s*[:=]\s*(?:['\"])?(?:INR\s*)?([\d,]+(?:\.\d+)?)",
        r"Portfolio\s*Value\s*[:=]\s*(?:INR|USD|EUR|GBP|₹|Rs\.?)?\s*([\d,]+(?:\.\d+)?)",
        r"portfolio_value\s*[:=]\s*(?:['\"])?(?:INR\s*)?([\d,]+(?:\.\d+)?)"], text)
    try:
        return float(m.group(1).replace(",", "")) if m else None
    except Exception:
        return None


ASSET_HEADER = re.compile(r"^\s*(?:portfolio\s+)?asset[\s_\-]*([A-Za-z]|\d+)(?![A-Za-z0-9])", re.IGNORECASE)
STOP_LINE = re.compile(r"^\s*(?:constraints?|scenario|stress|market\s+shocks?|horizon|institution|"
                       r"portfolio\s+amount|currency)\b", re.IGNORECASE)


def parse_sensitivities(block: str) -> dict[str, float]:
    out: dict[str, float] = {}
    for variable, aliases in VARIABLE_ALIASES.items():
        a = "(?:" + "|".join(aliases) + ")"
        m = find_first([rf"{a}\s+Sensitivity\s*[:=]\s*([+-]?\d+(?:\.\d+)?)",
                        rf"{a}\s*[-:]?\s*Sensitivity\s*[:=]\s*([+-]?\d+(?:\.\d+)?)",
                        rf"{a}\s*=\s*([+-]?\d+(?:\.\d+)?)"], block)
        if m:
            out[variable] = float(m.group(1))
    return out


def parse_assets(text: str) -> dict[str, dict[str, Any]]:
    assets: dict[str, dict[str, Any]] = {}
    lines = text.splitlines()
    starts = [(i, f"Asset_{m.group(1).upper()}") for i, ln in enumerate(lines) if (m := ASSET_HEADER.match(ln))]
    for k, (s, name) in enumerate(starts):
        limit = starts[k + 1][0] if k + 1 < len(starts) else len(lines)
        end = next((j for j in range(s + 1, limit) if STOP_LINE.match(lines[j])), limit)
        block = "\n".join(lines[s:end])
        m = find_first([r"Weight\s*[:=]\s*([+-]?\d+(?:\.\d+)?\s*%)", r"weight\s*[:=]\s*([+-]?\d+(?:\.\d+)?)",
                        r"portfolio_weight\s*[:=]\s*([+-]?\d+(?:\.\d+)?)"], block)
        weight = parse_percent_or_decimal(m.group(1)) if m else None
        sens = parse_sensitivities(block)
        if weight is not None or sens:
            assets[name] = {"weight": weight, "sensitivities": sens}

    # python-dict style:  "Asset_A": {"weight": 0.4, "sensitivities": {...}}
    for m in re.finditer(r"[\"']?(Asset[_\s\-]?[A-Za-z0-9]+)[\"']?\s*:\s*\{(?P<body>.*?)\}", text,
                         re.IGNORECASE | re.DOTALL):
        key = re.search(r"asset[\s_\-]*(\w+)", m.group(1), re.IGNORECASE)
        if not key:
            continue
        name, body = f"Asset_{key.group(1).upper()}", m.group("body")
        wm = re.search(r"['\"]?weight['\"]?\s*:\s*([-+]?\d+(?:\.\d+)?)", body, re.IGNORECASE)
        weight = parse_percent_or_decimal(wm.group(1)) if wm else None
        sens = parse_sensitivities(body)
        if weight is not None or sens:
            cur = assets.setdefault(name, {"weight": None, "sensitivities": {}})
            if weight is not None:
                cur["weight"] = weight
            cur["sensitivities"].update(sens)
    return assets


def parse_constraints(text: str) -> dict[str, float]:
    lo = hi = None
    m = find_first([r"min(?:imum)?\s*(?:weight)?\s*[:=]?\s*([+-]?\d+(?:\.\d+)?\s*%)",
                    r"min_weight\s*[:=]\s*([+-]?\d+(?:\.\d+)?)",
                    r"minimum_weight\s*[:=]\s*([+-]?\d+(?:\.\d+)?)"], text)
    if m:
        lo = parse_percent_or_decimal(m.group(1))
    m = find_first([r"max(?:imum)?\s*(?:weight)?\s*[:=]?\s*([+-]?\d+(?:\.\d+)?\s*%)",
                    r"max_weight\s*[:=]\s*([+-]?\d+(?:\.\d+)?)",
                    r"maximum_weight\s*[:=]\s*([+-]?\d+(?:\.\d+)?)"], text)
    if m:
        hi = parse_percent_or_decimal(m.group(1))
    m = re.search(r"Constraints.*?min(?:imum)?\s*[:=]?\s*([+-]?\d+(?:\.\d+)?\s*%).*?"
                  r"max(?:imum)?\s*[:=]?\s*([+-]?\d+(?:\.\d+)?\s*%)", text, re.IGNORECASE | re.DOTALL)
    if m:
        lo, hi = parse_percent_or_decimal(m.group(1)), parse_percent_or_decimal(m.group(2))
    return {"min_weight": 0.05 if lo is None else float(lo), "max_weight": 0.60 if hi is None else float(hi)}


def parse_scenario_name(text: str) -> str:
    m = find_first([r"Scenario\s*Name\s*[:=]\s*([^\n]+)", r"scenario_name\s*[:=]\s*[\"']([^\"']+)[\"']",
                    r"Scenario\s*[:=]\s*([^\n]+)", r"scenario\s*[:=]\s*[\"']([^\"']+)[\"']"], text)
    return m.group(1).strip().strip("\"'") if m else "Uploaded Stress Scenario"


def parse_horizon(text: str) -> int:
    m = find_first([r"Horizon\s*Days\s*[:=]\s*(\d+)", r"horizon_days\s*[:=]\s*(\d+)",
                    r"Horizon\s*[:=]\s*(\d+)\s*days?"], text)
    return max(1, int(m.group(1))) if m else 30


def parse_shocks(text: str) -> dict[str, float]:
    shocks: dict[str, float] = {}
    for variable, aliases in VARIABLE_ALIASES.items():
        a = "(?:" + "|".join(aliases) + ")"
        m = find_first([rf"{a}\s+Shock\s*[:=]\s*([+-]?\d+(?:\.\d+)?\s*(?:%|percentage\s+points?)?)",
                        rf"{a}\s*Shock\s*[:=]\s*([+-]?\d+(?:\.\d+)?)",
                        rf"{a}\s*[:=]\s*([+-]?\d+(?:\.\d+)?\s*%)"], text)
        if m:
            v = parse_shock_value(m.group(1), variable)
            if v is not None:
                shocks[variable] = v
    return shocks


def parse_uploaded_config(text: str) -> dict[str, Any]:
    text = clean_text(text)
    return {
        "type": parse_institution_type(text),
        "portfolio_amount": parse_portfolio_amount(text),
        "currency": parse_currency(text),
        "portfolio_assets": parse_assets(text),
        "constraints": parse_constraints(text),
        "scenario": {"name": parse_scenario_name(text), "horizon_days": parse_horizon(text),
                     "shocks": parse_shocks(text)},
        "raw_text": text,
    }


# ============================================================
# VALIDATION / AGENT INPUTS
# ============================================================
def validate_config(config: dict[str, Any] | None) -> tuple[bool, list[str]]:
    if not config:
        return False, ["No configuration was parsed."]
    errors: list[str] = []
    if "bank" in str(config.get("type", "")).lower():
        errors.append("Banking institutions are not connected to the dynamic pipeline yet "
                      "(Stock Market only).")
    amount = config.get("portfolio_amount")
    if amount is None:
        errors.append("Portfolio Amount is missing.")
    elif amount <= 0:
        errors.append("Portfolio Amount must be greater than zero.")

    assets = config.get("portfolio_assets", {})
    c = config.get("constraints", {})
    lo, hi = c.get("min_weight"), c.get("max_weight")
    if not assets:
        errors.append("No portfolio assets were detected.")
    if len(assets) == 1:
        errors.append("At least two portfolio assets are required.")

    total = 0.0
    for name, a in assets.items():
        w = a.get("weight")
        if w is None:
            errors.append(f"{name}: weight is missing.")
        else:
            total += float(w)
            if lo is not None and hi is not None and not (lo - 1e-9 <= w <= hi + 1e-9):
                errors.append(f"{name}: weight {w:.2%} is outside the constraints ({lo:.0%}–{hi:.0%}).")
        if not a.get("sensitivities"):
            errors.append(f"{name}: no sensitivities detected.")
    if assets and abs(total - 1.0) > 0.02:
        errors.append(f"Portfolio weights sum to {total * 100:.2f}%. They should total about 100%.")

    if lo is None or hi is None:
        errors.append("Portfolio weight constraints are missing.")
    else:
        if lo < 0:
            errors.append("Minimum weight cannot be negative.")
        if hi > 1:
            errors.append("Maximum weight cannot exceed 100%.")
        if lo > hi:
            errors.append("Minimum weight cannot exceed maximum weight.")
        if assets and len(assets) * lo > 1 + 1e-9:
            errors.append("Constraints are infeasible: assets × minimum weight exceeds 100%.")
        if assets and len(assets) * hi < 1 - 1e-9:
            errors.append("Constraints are infeasible: assets × maximum weight is below 100%.")

    sc = config.get("scenario", {})
    if not sc.get("horizon_days") or sc["horizon_days"] <= 0:
        errors.append("Scenario horizon must be greater than zero.")
    for v in RISK_VARIABLES:
        if v not in sc.get("shocks", {}):
            errors.append(f"Stress shock for {v} is missing.")
    return not errors, errors


def normalized_weights(config: dict[str, Any]) -> tuple[list[str], np.ndarray]:
    names = list(config["portfolio_assets"].keys())
    w = np.array([float(config["portfolio_assets"][n]["weight"]) for n in names])
    return names, w / w.sum()


def build_institution_config(config: dict[str, Any]) -> dict[str, Any]:
    names, w = normalized_weights(config)
    assets = {n: {"weight": float(wi),
                  "sensitivities": {normalize_key(k): float(v)
                                    for k, v in config["portfolio_assets"][n]["sensitivities"].items()}}
              for n, wi in zip(names, w)}
    return {"type": config["type"], "portfolio_amount": float(config["portfolio_amount"]),
            "currency": config["currency"], "portfolio_assets": assets,
            "constraints": dict(config["constraints"]), "project_root": str(PROJECT_ROOT)}


def build_agent_scenario(config: dict[str, Any]) -> dict[str, Any]:
    """The agent treats |value| > 1 as a percentage, so shocks above +/-100% are passed as percent numbers."""
    sc = config["scenario"]
    out: dict[str, Any] = {"name": sc["name"], "horizon_days": int(sc["horizon_days"])}
    for k, v in sc["shocks"].items():
        k = normalize_key(k)
        out[k] = float(v) if (k == "INDIA_10Y_YIELD" or abs(v) <= 1.0) else float(v) * 100.0
    return out


# ============================================================
# AGENTS
# ============================================================
def _import_first(candidates: list[tuple[str, str]], what: str):
    errors = []
    for module_name, class_name in candidates:
        try:
            return getattr(importlib.import_module(module_name), class_name)
        except Exception as exc:
            errors.append(f"{module_name}.{class_name}: {exc}")
    raise ImportError(f"Could not import {what}.\n\n" + "\n".join(errors))


@st.cache_resource(show_spinner="Loading TimeGAN model...")
def get_simulation_agent():
    cls = _import_first([("agents.simulation_agent", "SimulationAgent"),
                         ("agent.simulation_agent", "SimulationAgent"),
                         ("src.agents.simulation_agent", "SimulationAgent"),
                         ("simulation_agent", "SimulationAgent")], "SimulationAgent")
    return cls()


def get_optimizer():
    cls = _import_first([("agents.portfolio_optimization_agent_v2", "PortfolioOptimizationAgentV2"),
                         ("agents.portfolio_optimization_agent", "PortfolioOptimizationAgentV2"),
                         ("agent.portfolio_optimization_agent_v2", "PortfolioOptimizationAgentV2"),
                         ("src.agents.portfolio_optimization_agent_v2", "PortfolioOptimizationAgentV2"),
                         ("portfolio_optimization_agent_v2", "PortfolioOptimizationAgentV2")],
                        "PortfolioOptimizationAgentV2")
    return cls()


def ensure_raw_array(raw: Any, name: str) -> np.ndarray:
    arr = np.asarray(raw, dtype=float)
    if arr.ndim != 3 or arr.shape[2] != 5:
        raise ValueError(f"{name} must have shape (scenarios, days, 5). Received {arr.shape}.")
    return arr


def run_simulation(config: dict[str, Any]) -> dict[str, Any]:
    sim = get_simulation_agent().run(
        scenario=build_agent_scenario(config), num_scenarios=NUM_SCENARIOS,
        horizon_days=int(config["scenario"]["horizon_days"]), seed=DEFAULT_SEED,
        portfolio_amount=float(config["portfolio_amount"]))
    if not isinstance(sim, dict) or "baseline_raw" not in sim or "stressed_raw" not in sim:
        raise RuntimeError("SimulationAgent did not return baseline_raw / stressed_raw.")
    return sim


def run_optimization(config: dict[str, Any], sim: dict[str, Any]) -> dict[str, Any]:
    result = get_optimizer().optimize(
        institution=build_institution_config(config),
        baseline_raw=ensure_raw_array(sim["baseline_raw"], "baseline_raw"),
        stressed_raw=ensure_raw_array(sim["stressed_raw"], "stressed_raw"),
        save_output=True)
    if not isinstance(result, dict):
        raise RuntimeError("PortfolioOptimizationAgentV2 returned an unexpected result.")
    return result


def extract_weights_from_result(result: dict[str, Any], names: list[str]) -> np.ndarray | None:
    for obj in (result.get("optimized_portfolio"), result.get("optimization_result"), result):
        if not isinstance(obj, dict):
            continue
        for key in ("weights", "optimized_weights", "portfolio_weights"):
            w = obj.get(key)
            if isinstance(w, dict) and all(n in w for n in names):
                return np.array([float(w[n]) for n in names])
            if isinstance(w, (list, tuple, np.ndarray)) and len(w) == len(names):
                return np.asarray(w, dtype=float).reshape(-1)
    return None


# ============================================================
# RISK ENGINE
# ============================================================
def portfolio_returns(raw: np.ndarray, config: dict[str, Any], weights: np.ndarray) -> np.ndarray:
    names = list(config["portfolio_assets"].keys())
    matrix = np.zeros((len(names), len(STOCK_FEATURES)))
    for i, n in enumerate(names):
        for var, s in config["portfolio_assets"][n]["sensitivities"].items():
            idx = RISK_TO_INDEX.get(normalize_key(var))
            if idx is not None:
                matrix[i, idx] = float(s)
    return np.einsum("sda,a->sd", np.einsum("sdf,af->sda", raw, matrix), weights)


def risk_metrics(daily: np.ndarray, value: float, conf: float) -> dict[str, float]:
    """Cumulative-path metrics (same definition the optimizer minimizes)."""
    r = np.clip(np.nan_to_num(np.asarray(daily, dtype=float), nan=0.0, posinf=0.0, neginf=0.0), -0.999999, None)
    cum = np.cumprod(1.0 + r, axis=1)
    loss = np.maximum(-(cum[:, -1] - 1.0), 0.0)
    var = float(np.quantile(loss, conf))
    tail = loss[loss >= var]
    peak = np.maximum.accumulate(cum, axis=1)
    vals = {"VaR": var, "ES": float(tail.mean() if len(tail) else var), "Expected Loss": float(loss.mean()),
            "Max Drawdown": float(((peak - cum) / np.maximum(peak, 1e-12)).max()), "Worst Loss": float(loss.max())}
    out = dict(vals)
    out.update({f"{k} INR": v * value for k, v in vals.items()})
    return out


def metrics_table(m: dict[str, float], conf: float, cur: str) -> pd.DataFrame:
    return pd.DataFrame([{"Risk Metric": label(k, conf), "Percentage": pct(m[k]),
                          "Monetary Exposure": money(m[k + " INR"], cur)} for k in RISK_KEYS])


def day_by_day(daily: np.ndarray, value: float, conf: float, cur: str) -> pd.DataFrame:
    loss = np.maximum(-(np.cumprod(1 + np.clip(daily, -0.999999, None), axis=1) - 1), 0)
    rows = []
    for d in range(loss.shape[1]):
        v = loss[:, d]
        var = np.quantile(v, conf)
        tail = v[v >= var]
        rows.append({"Day": d + 1, label("VaR", conf): pct(var),
                     label("ES", conf): pct(tail.mean() if len(tail) else var),
                     "Expected Loss": pct(v.mean()), "Worst Loss": pct(v.max()),
                     "VaR Exposure": money(var * value, cur), "Expected Loss Exposure": money(v.mean() * value, cur)})
    return pd.DataFrame(rows)


# ---- NEW: day-to-day (single-day) loss -----------------------------------
def daily_loss_frame(daily: np.ndarray, value: float, conf: float) -> pd.DataFrame:
    """Numeric per-day loss statistics across all scenarios (single-day loss, not cumulative)."""
    r = np.clip(np.nan_to_num(np.asarray(daily, dtype=float)), -0.999999, None)
    day_loss = -r                                          # positive = loss on that day
    cum_loss = np.maximum(-(np.cumprod(1.0 + r, axis=1) - 1.0), 0.0)
    rows = []
    for d in range(r.shape[1]):
        dl, cl = day_loss[:, d], cum_loss[:, d]
        rows.append({
            "Day": d + 1,
            "Avg Daily Return": r[:, d].mean(),
            "Avg Daily Loss": np.maximum(dl, 0).mean(),
            "Daily VaR": max(float(np.quantile(dl, conf)), 0.0),
            "Worst Daily Loss": max(float(dl.max()), 0.0),
            "Cum Expected Loss": cl.mean(),
            "Cum P05": np.quantile(cl, 0.05),
            "Cum P95": np.quantile(cl, 0.95),
            "Cum VaR": np.quantile(cl, conf),
            "Cum Worst": cl.max(),
            "Daily Loss INR": np.maximum(dl, 0).mean() * value,
            "Cum Loss INR": cl.mean() * value,
        })
    return pd.DataFrame(rows)


def daily_loss_table(df: pd.DataFrame, conf: float, cur: str) -> pd.DataFrame:
    return pd.DataFrame({
        "Day": df["Day"],
        "Avg Daily Return": df["Avg Daily Return"].map(pct),
        "Avg Daily Loss": df["Avg Daily Loss"].map(pct),
        f"Daily VaR {conf:.0%}": df["Daily VaR"].map(pct),
        "Worst Daily Loss": df["Worst Daily Loss"].map(pct),
        "Avg Daily Loss (Amount)": df["Daily Loss INR"].map(lambda v: money(v, cur)),
        "Cumulative Expected Loss": df["Cum Expected Loss"].map(pct),
        "Cumulative Loss (Amount)": df["Cum Loss INR"].map(lambda v: money(v, cur)),
    })


def calculate_restress(config: dict[str, Any], sim: dict[str, Any], opt: dict[str, Any]) -> dict[str, Any]:
    names, base_w = normalized_weights(config)
    opt_w = extract_weights_from_result(opt, names)
    if opt_w is None:
        raise RuntimeError("Could not extract optimized weights from the optimizer result.")
    raw = ensure_raw_array(sim["stressed_raw"], "stressed_raw")
    return {"assets": names, "base_w": base_w, "opt_w": opt_w,
            "base_daily": portfolio_returns(raw, config, base_w),
            "opt_daily": portfolio_returns(raw, config, opt_w)}


def restress_checks(r: dict[str, Any], bm: dict, om: dict, config: dict[str, Any]) -> list[tuple[str, bool]]:
    lo, hi = config["constraints"]["min_weight"], config["constraints"]["max_weight"]
    w = r["opt_w"]
    same = r["base_daily"].shape == r["opt_daily"].shape
    return [
        ("Same TimeGAN scenario paths reused", same),
        ("No new synthetic data generated", True),   # true by design: no generator is called in re-stress
        ("Optimized weights sum to 100%", bool(abs(w.sum() - 1) < 1e-3)),
        (f"Weights remain within {lo:.0%}–{hi:.0%} bounds", bool(np.all((w >= lo - 1e-6) & (w <= hi + 1e-6)))),
        ("Expected loss improved", om["Expected Loss"] < bm["Expected Loss"]),
        ("Expected Shortfall (ES) improved", om["ES"] < bm["ES"]),
        ("Maximum drawdown improved", om["Max Drawdown"] < bm["Max Drawdown"]),
        ("Baseline and optimized portfolios evaluated on identical scenarios", same),
    ]


# ============================================================
# CHARTS
# ============================================================
def drawdown_chart(daily: np.ndarray) -> go.Figure:
    cum = np.cumprod(1 + np.clip(daily, -0.999999, None), axis=1)
    peak = np.maximum.accumulate(cum, axis=1)
    dd = (peak - cum) / np.maximum(peak, 1e-12)
    days = np.arange(1, daily.shape[1] + 1)
    fig = go.Figure()
    fig.add_trace(go.Scatter(x=days, y=dd.mean(axis=0) * 100, mode="lines", name="Average Drawdown",
                             line=dict(color=BLUE, width=3), fill="tozeroy", fillcolor="rgba(37,99,235,.12)"))
    fig.add_trace(go.Scatter(x=days, y=dd.max(axis=0) * 100, mode="lines", name="Worst Drawdown",
                             line=dict(color=RED, width=3)))
    fig.update_layout(title=f"{daily.shape[1]}-Day Portfolio Drawdown", xaxis_title="Day", yaxis_title="Drawdown (%)")
    return style_fig(fig)


def compare_chart(bm: dict, om: dict, conf: float) -> go.Figure:
    names = [label(k, conf) for k in RISK_KEYS]
    fig = go.Figure([go.Bar(x=names, y=[bm[k] * 100 for k in RISK_KEYS], name="Baseline", marker_color=MUTED),
                     go.Bar(x=names, y=[om[k] * 100 for k in RISK_KEYS], name="Optimized", marker_color=GREEN)])
    fig.update_layout(title="Baseline vs Optimized Risk", xaxis_title="Risk Metric", yaxis_title="Risk (%)",
                      barmode="group")
    return style_fig(fig)


def weights_chart(names: list[str], base: np.ndarray, opt: np.ndarray) -> go.Figure:
    fig = go.Figure([go.Bar(x=names, y=base * 100, name="Baseline", marker_color=MUTED),
                     go.Bar(x=names, y=opt * 100, name="Optimized", marker_color=BLUE)])
    fig.update_layout(title="Portfolio Allocation", xaxis_title="Asset", yaxis_title="Weight (%)", barmode="group")
    return style_fig(fig)


# ---- NEW: day-to-day loss charts -----------------------------------------
def daily_loss_chart(df: pd.DataFrame, conf: float) -> go.Figure:
    fig = go.Figure()
    fig.add_trace(go.Bar(x=df["Day"], y=df["Avg Daily Loss"] * 100, name="Average daily loss",
                         marker_color=BLUE, opacity=0.85))
    fig.add_trace(go.Scatter(x=df["Day"], y=df["Daily VaR"] * 100, mode="lines", name=f"Daily VaR {conf:.0%}",
                             line=dict(color=AMBER, width=3)))
    fig.add_trace(go.Scatter(x=df["Day"], y=df["Worst Daily Loss"] * 100, mode="lines", name="Worst daily loss",
                             line=dict(color=RED, width=2, dash="dot")))
    fig.update_layout(title="Day-to-Day Portfolio Loss (single-day loss across scenarios)",
                      xaxis_title="Day", yaxis_title="Loss on the day (%)")
    return style_fig(fig)


def cumulative_loss_chart(df: pd.DataFrame, conf: float) -> go.Figure:
    fig = go.Figure()
    fig.add_trace(go.Scatter(x=df["Day"], y=df["Cum P95"] * 100, mode="lines", line=dict(width=0),
                             showlegend=False, hoverinfo="skip"))
    fig.add_trace(go.Scatter(x=df["Day"], y=df["Cum P05"] * 100, mode="lines", line=dict(width=0),
                             fill="tonexty", fillcolor="rgba(37,99,235,.15)", name="5th–95th percentile band"))
    fig.add_trace(go.Scatter(x=df["Day"], y=df["Cum Expected Loss"] * 100, mode="lines",
                             name="Expected cumulative loss", line=dict(color=BLUE, width=3)))
    fig.add_trace(go.Scatter(x=df["Day"], y=df["Cum VaR"] * 100, mode="lines", name=f"Cumulative VaR {conf:.0%}",
                             line=dict(color=AMBER, width=3)))
    fig.add_trace(go.Scatter(x=df["Day"], y=df["Cum Worst"] * 100, mode="lines", name="Worst cumulative loss",
                             line=dict(color=RED, width=2, dash="dot")))
    fig.update_layout(title="Cumulative Portfolio Loss Over Time", xaxis_title="Day",
                      yaxis_title="Cumulative loss (%)")
    return style_fig(fig)


def compare_daily_loss_chart(base_daily: np.ndarray, opt_daily: np.ndarray, value: float, conf: float) -> go.Figure:
    b, o = daily_loss_frame(base_daily, value, conf), daily_loss_frame(opt_daily, value, conf)
    fig = go.Figure()
    fig.add_trace(go.Scatter(x=b["Day"], y=b["Cum Expected Loss"] * 100, mode="lines", name="Baseline",
                             line=dict(color=MUTED, width=3)))
    fig.add_trace(go.Scatter(x=o["Day"], y=o["Cum Expected Loss"] * 100, mode="lines", name="Optimized",
                             line=dict(color=GREEN, width=3)))
    fig.update_layout(title="Day-by-Day Expected Loss: Baseline vs Optimized", xaxis_title="Day",
                      yaxis_title="Cumulative expected loss (%)")
    return style_fig(fig)


# ============================================================
# REPORT
# ============================================================
def build_report_html(config, sim, r, bm, om, conf, checks) -> str:
    cur = config["currency"]
    tbl = lambda head, rows: ("<table><tr>" + "".join(f"<th>{h}</th>" for h in head) + "</tr>" +
                              "".join("<tr>" + "".join(f"<td>{c}</td>" for c in row) + "</tr>" for row in rows) +
                              "</table>")
    shocks = tbl(["Variable", "Shock"], [[v, shock_display(v, config["scenario"]["shocks"].get(v))]
                                         for v in RISK_VARIABLES])
    risk = tbl(["Metric", "Baseline portfolio", "Optimized portfolio"],
               [[label(k, conf), f"{pct(bm[k])} ({money(bm[k + ' INR'], cur)})",
                 f"{pct(om[k])} ({money(om[k + ' INR'], cur)})"] for k in RISK_KEYS])
    wts = tbl(["Asset", "Baseline", "Optimized"],
              [[a, pct(b), pct(o)] for a, b, o in zip(r["assets"], r["base_w"], r["opt_w"])])
    verdict = "PASSED" if all(ok for _, ok in checks) else "FAILED"
    chk = tbl(["Result", "Check"], [["PASS" if ok else "FAIL", t] for t, ok in checks])

    # day-to-day loss section (baseline vs optimized)
    value = float(config["portfolio_amount"])
    bdf, odf = daily_loss_frame(r["base_daily"], value, conf), daily_loss_frame(r["opt_daily"], value, conf)
    daily_rows = [[int(b["Day"]), pct(b["Avg Daily Loss"]), pct(o["Avg Daily Loss"]),
                   pct(b["Cum Expected Loss"]), pct(o["Cum Expected Loss"])]
                  for (_, b), (_, o) in zip(bdf.iterrows(), odf.iterrows())]
    daily_tbl = tbl(["Day", "Baseline avg daily loss", "Optimized avg daily loss",
                     "Baseline cumulative loss", "Optimized cumulative loss"], daily_rows)

    raw = sim["stressed_raw"]
    return f"""<!DOCTYPE html><html><head><meta charset="UTF-8"><title>MacroStress-GAN Risk Report</title>
<style>body{{font-family:Arial,sans-serif;margin:40px;color:#0F172A}}h1{{color:#0F2A47}}
h2{{color:#1E4E8C;margin-top:30px;border-left:5px solid #0D9488;padding-left:10px}}
table{{width:100%;border-collapse:collapse;margin-top:15px}}th,td{{border:1px solid #CBD5E1;padding:10px;text-align:left}}
th{{background:#0F2A47;color:#fff}}tr:nth-child(even) td{{background:#F1F5F9}}
.summary{{padding:15px;background:#EFF6FF;border-left:5px solid #2563EB;border-radius:8px}}</style></head><body>
<h1>MacroStress-GAN Risk Report</h1>
<div class="summary"><p><strong>Institution Type:</strong> {config['type']}</p>
<p><strong>Portfolio Value:</strong> {money(config['portfolio_amount'], cur)}</p>
<p><strong>Scenario:</strong> {config['scenario']['name']} · {raw.shape[1]} days · {raw.shape[0]:,} paths</p>
<p><strong>Confidence:</strong> {conf:.0%} · <strong>Risk Rating:</strong> {risk_rating(bm['Expected Loss'])}</p>
<p><strong>Generated:</strong> {datetime.now():%d %b %Y %H:%M} · source: uploaded document only</p></div>
<h2>Stress Shocks</h2>{shocks}<h2>Risk Metrics</h2>{risk}<h2>Portfolio Allocation</h2>{wts}
<h2>Day-to-Day Loss</h2>{daily_tbl}
<h2>Re-Stress Verification: {verdict}</h2>{chk}
<h2>Methodology</h2><p>TimeGAN generates baseline paths; the uploaded scenario shocks are injected and validated.
Portfolio risk is evaluated on the stressed paths. The optimizer minimizes expected stressed loss under the uploaded
weight bounds, and the result is re-evaluated on the same paths (in-sample check). For risk analysis only.</p>
</body></html>"""


def export_json(config, r, bm, om, conf, checks) -> bytes:
    payload = {"application": APP_NAME, "version": APP_VERSION, "confidence": conf,
               "config": {k: v for k, v in config.items() if k != "raw_text"},
               "weights": {"baseline": dict(zip(r["assets"], map(float, r["base_w"]))),
                           "optimized": dict(zip(r["assets"], map(float, r["opt_w"])))},
               "baseline_metrics": bm, "optimized_metrics": om,
               "restress_checks": [{"check": t, "passed": bool(ok)} for t, ok in checks]}
    return json.dumps(payload, indent=2, default=float).encode("utf-8")


# ============================================================
# SIDEBAR / LANDING
# ============================================================
def render_sidebar() -> tuple[bool, float]:
    with st.sidebar:
        st.markdown("## MacroStress-GAN")
        st.caption("Multi-Agent Generative AI · Financial Stress Testing")
        st.divider()
        up = st.file_uploader("Upload institution document", type=["pdf", "docx", "txt", "py", "md"],
                              help="Institution Type, Portfolio Amount, Assets, Weights, Sensitivities, "
                                   "Constraints and Stress Scenario.")
        st.caption("200MB per file • DOCX, PDF, TXT")
        if up is not None:
            key = f"{up.name}_{up.size}"
            if st.session_state.upload_key != key:
                try:
                    st.session_state.config = parse_uploaded_config(read_uploaded_file(up))
                    st.session_state.upload_key, st.session_state.upload_name = key, up.name
                    for k in ("simulation", "optimization", "restress", "pipeline_error"):
                        st.session_state[k] = None          # a new upload invalidates every earlier run
                except Exception as exc:
                    st.session_state.config = None
                    st.session_state.pipeline_error = f"Could not read uploaded file: {exc}"
        st.divider()
        st.markdown("### Settings")
        conf = st.slider("Confidence level", 0.90, 0.99, 0.95, 0.01, format="%.2f")
        st.caption(f"Risk confidence: {conf:.0%}")

        cfg = st.session_state.config
        if cfg:
            uk = st.session_state.upload_key
            st.divider()
            st.subheader("Portfolio Controls")
            st.number_input("Portfolio Value (₹)", min_value=0.0, value=float(cfg["portfolio_amount"] or 0.0),
                            format="%.0f", disabled=True, key=f"pv_{uk}")
            st.caption("Baseline allocation (from document)")
            for name, a in cfg["portfolio_assets"].items():
                w = min(max(float(a.get("weight") or 0.0), 0.0), 1.0)
                st.slider(name.replace("_", " "), 0.0, 1.0, w, 0.01, format="%.3f",
                          disabled=True, key=f"w_{name}_{uk}")
            st.caption("These values come from the uploaded document. Edit the document to change them.")
            st.divider()
        run = st.button("▶ Run Dynamic Stress Test", type="primary", width="stretch", disabled=cfg is None)
        if st.button("Clear current run", width="stretch"):
            for k, v in SESSION_DEFAULTS.items():
                st.session_state[k] = v
            st.rerun()
        st.caption(f"Version {APP_VERSION}")
    return run, conf


def landing() -> None:
    st.markdown('<div class="hero-box"><h1>MacroStress-GAN Risk Platform</h1>'
                '<p>Generative-AI stress testing, risk analytics, optimization and re-stress verification.</p></div>',
                unsafe_allow_html=True)
    if st.session_state.pipeline_error:
        st.error(st.session_state.pipeline_error)
    st.info("Upload an institution document to begin. The platform reads the portfolio, assets, constraints and "
            "stress scenario directly from it, then runs TimeGAN, risk analytics, optimization and re-stress. "
            "No old results are reused.")
    c1, c2 = st.columns(2)
    c1.markdown(f'<div class="info-card" style="border-top:4px solid {BLUE};"><h3>📈 Stock Market</h3>'
                '<p>Institution Type: Stock Market</p><p>NIFTY50 · CRUDE_OIL · USD_INR · INDIA_VIX · INDIA_10Y_YIELD</p></div>',
                unsafe_allow_html=True)
    c2.markdown(f'<div class="info-card" style="border-top:4px solid {TEAL};"><h3>🏦 Banking</h3>'
                '<p>Institution Type: Banking</p><p>BANK_NIFTY · REPO_RATE · CPI_INFLATION · CREDIT_GROWTH · USD_INR</p>'
                '<p><i>Not connected to the dynamic pipeline yet.</i></p></div>', unsafe_allow_html=True)
    st.markdown('<div class="section-title">Platform Workflow</div>', unsafe_allow_html=True)
    steps = ["Document ingestion", "Institution detection", "TimeGAN scenario generation",
             "1,000 × N-day stress simulations", "Day-by-day portfolio risk",
             "VaR / ES / Expected Loss / Drawdown", "Portfolio optimization",
             "Re-stress verification", "Final risk report"]
    cols = st.columns(3)
    for i, s in enumerate(steps):
        cols[i % 3].markdown(f'<div class="step"><div class="step-no">{i + 1}</div>{s}</div>', unsafe_allow_html=True)


def hero(institution: str) -> None:
    st.markdown(f'<div class="hero-box"><h1>MacroStress-GAN Risk Platform</h1>'
                f'<p>{institution} · TimeGAN Stress Testing · Portfolio Risk Analytics</p></div>',
                unsafe_allow_html=True)


# ============================================================
# TABS
# ============================================================
def risk_kpis(m: dict[str, float], conf: float, cur: str) -> None:
    cols = st.columns(5)
    specs = [("VaR", "amber"), ("ES", "orange"), ("Expected Loss", "blue"), ("Max Drawdown", "red"),
             ("Worst Loss", "red")]
    for col, (key, accent) in zip(cols, specs):
        with col:
            kpi(label(key, conf), pct(m[key]), money(m[key + " INR"], cur), accent)


def tab_overview(config, sim, r, bm, conf):
    cur, raw = config["currency"], sim["stressed_raw"]
    value = float(config["portfolio_amount"])
    title("Risk Overview", "Portfolio-level stress analytics from generated TimeGAN scenarios.")
    risk_kpis(bm, conf, cur)
    st.markdown("")
    c1, c2 = st.columns([1.4, 1])
    with c1:
        st.plotly_chart(drawdown_chart(r["base_daily"]), width="stretch", key="overview_drawdown_chart")
        st.plotly_chart(daily_loss_chart(daily_loss_frame(r["base_daily"], value, conf), conf),
                        width="stretch", key="overview_daily_loss_chart")
    with c2:
        st.subheader("Simulation Summary")
        st.dataframe(pd.DataFrame({
            "Parameter": ["Institution", "Scenarios", "Days / Scenario", "Features", "Portfolio Value", "Risk Rating"],
            "Value": [config["type"], f"{raw.shape[0]:,}", raw.shape[1], raw.shape[2],
                      money(config["portfolio_amount"], cur), risk_rating(bm["Expected Loss"])]}),
            width="stretch", hide_index=True)
    title("Portfolio Exposure", "Baseline allocation used by the risk engine (from the uploaded document).")
    st.dataframe(pd.DataFrame({"Asset": r["assets"], "Weight": [pct(w) for w in r["base_w"]]}),
                 width="stretch", hide_index=True)


def tab_risk(config, r, bm, conf):
    cur, value = config["currency"], float(config["portfolio_amount"])
    title("Risk Analytics", "Day-by-day stress risk calculated across all generated scenarios.")
    risk_kpis(bm, conf, cur)
    st.markdown("")
    st.plotly_chart(drawdown_chart(r["base_daily"]), width="stretch", key="risk_drawdown_chart")
    title("Risk Metrics", "Percentage loss and corresponding monetary exposure.")
    st.dataframe(metrics_table(bm, conf, cur), width="stretch", hide_index=True)
    title("Day-by-Day Risk", "Each day shows the distribution of cumulative portfolio losses across all scenarios.")
    table = day_by_day(r["base_daily"], value, conf, cur)
    st.dataframe(table, width="stretch", hide_index=True)
    st.download_button("Download Day-by-Day Risk CSV", data=table.to_csv(index=False).encode("utf-8"),
                       file_name="macro_stress_day_by_day_risk.csv", mime="text/csv", key="download_day_by_day_csv")

    # ---- NEW: day-to-day loss ----
    title("Day-to-Day Loss", "Loss incurred on each individual day, plus how it accumulates over the horizon.")
    dl = daily_loss_frame(r["base_daily"], value, conf)
    g1, g2 = st.columns(2)
    with g1:
        st.plotly_chart(daily_loss_chart(dl, conf), width="stretch", key="risk_daily_loss_chart")
    with g2:
        st.plotly_chart(cumulative_loss_chart(dl, conf), width="stretch", key="risk_cum_loss_chart")
    dl_table = daily_loss_table(dl, conf, cur)
    st.dataframe(dl_table, width="stretch", hide_index=True)
    st.download_button("Download Day-to-Day Loss CSV", data=dl_table.to_csv(index=False).encode("utf-8"),
                       file_name="macro_stress_day_to_day_loss.csv", mime="text/csv", key="download_daily_loss_csv")


def tab_optimization(config, opt, r, bm, om, conf):
    cur = config["currency"]
    title("Portfolio Optimization", "Compare baseline allocation with the optimized allocation.")
    c1, c2 = st.columns(2)
    with c1:
        st.subheader("Baseline Allocation")
        st.dataframe(pd.DataFrame({"Asset": r["assets"], "Weight": [pct(v) for v in r["base_w"]]}),
                     width="stretch", hide_index=True)
    with c2:
        st.subheader("Optimized Allocation")
        st.dataframe(pd.DataFrame({"Asset": r["assets"], "Weight": [pct(v) for v in r["opt_w"]]}),
                     width="stretch", hide_index=True)
    st.plotly_chart(weights_chart(r["assets"], r["base_w"], r["opt_w"]), width="stretch",
                    key="optimization_weights_chart")
    st.plotly_chart(compare_chart(bm, om, conf), width="stretch", key="optimization_risk_chart")
    st.subheader("Optimization Impact")
    b, o = bm["Expected Loss INR"], om["Expected Loss INR"]
    cols = st.columns(3)
    with cols[0]:
        kpi("Baseline Expected Loss", money(b, cur), pct(bm["Expected Loss"]), "amber")
    with cols[1]:
        kpi("Optimized Expected Loss", money(o, cur), pct(om["Expected Loss"]), "green")
    with cols[2]:
        kpi("Loss Reduction", money(b - o, cur), pct((b - o) / b if b > 0 else 0), "teal")
    st.subheader("Optimized Risk Metrics")
    st.dataframe(metrics_table(om, conf, cur), width="stretch", hide_index=True)
    res, meta = opt.get("optimization_result", {}), opt.get("optimizer", {})
    if res or meta:
        with st.expander("Optimizer details"):
            st.json({"selected_method": res.get("selected_method"), "candidate_count": res.get("candidate_count"),
                     "slsqp_success": meta.get("slsqp_success"), "validation": opt.get("validation")})


def tab_restress(config, r, bm, om, conf, checks):
    cur = config["currency"]
    title("Re-Stress Verification",
          "The optimized portfolio is re-evaluated using the same TimeGAN-generated stress scenario paths. "
          "No new synthetic scenarios are generated. This verifies whether the optimized allocation produces "
          "lower modeled risk under the original stress conditions.")
    b, o = bm["Expected Loss INR"], om["Expected Loss INR"]
    red_pct = (b - o) / b if b > 0 else 0
    cols = st.columns(4)
    with cols[0]:
        kpi("Baseline Loss", money(b, cur), pct(bm["Expected Loss"]), "amber")
    with cols[1]:
        kpi("Optimized Loss", money(o, cur), pct(om["Expected Loss"]), "green")
    with cols[2]:
        kpi("Reduction", money(b - o, cur), pct(red_pct), "teal")
    with cols[3]:
        rating = risk_rating(om["Expected Loss"])
        kpi("Risk Rating", rating, "After re-stress", RATING_ACCENT[rating])
    st.markdown("")
    for text, ok in checks:
        st.markdown(f'<div class="check">{"✅" if ok else "❌"} {text}</div>', unsafe_allow_html=True)
    if all(ok for _, ok in checks):
        st.success("✅ VERIFICATION PASSED. The optimized allocation shows lower modeled risk on the original "
                   "stress scenarios.")
    else:
        st.error("Verification has failures. Review the flagged checks before adopting this allocation.")
    st.plotly_chart(compare_chart(bm, om, conf), width="stretch", key="restress_risk_chart")
    # ---- NEW: day-by-day baseline vs optimized ----
    st.plotly_chart(compare_daily_loss_chart(r["base_daily"], r["opt_daily"], float(config["portfolio_amount"]), conf),
                    width="stretch", key="restress_daily_loss_chart")
    st.subheader("Re-Stress Risk Metrics")
    st.dataframe(pd.DataFrame({"Metric": [label(k, conf) for k in RISK_KEYS],
                               "Baseline": [pct(bm[k]) for k in RISK_KEYS],
                               "Optimized": [pct(om[k]) for k in RISK_KEYS]}),
                 width="stretch", hide_index=True)


def tab_data(sim):
    title("Synthetic Scenario Data", "TimeGAN-generated scenarios used by the risk engine (current run).")
    raw = ensure_raw_array(sim["stressed_raw"], "stressed_raw")
    s, d, f = raw.shape
    df = pd.DataFrame(raw.reshape(-1, f), columns=STOCK_FEATURES)
    df.insert(0, "scenario", np.repeat(np.arange(s), d))
    df.insert(1, "day", np.tile(np.arange(d), s))
    cols = st.columns(4)
    for col, (lbl, val, acc) in zip(cols, [("Rows", f"{len(df):,}", "blue"), ("Scenarios", f"{s:,}", "teal"),
                                           ("Days", str(d), "green"), ("Features", str(f), "amber")]):
        with col:
            kpi(lbl, val, "", acc)
    st.markdown("")
    feats = sim.get("validation", {}).get("features")
    if feats:
        with st.expander("Scenario propagation check (baseline vs stressed means)"):
            st.dataframe(pd.DataFrame([{"Feature": k, "Baseline mean": f"{v['baseline_mean']:+.6f}",
                                        "Stressed mean": f"{v['stressed_mean']:+.6f}",
                                        "Delta": f"{v['actual_delta']:+.6f}", "Expected": f"{v['expected_delta']:+.6f}",
                                        "Passed": v["passed"]} for k, v in feats.items()]),
                         width="stretch", hide_index=True)
    st.subheader("Columns")
    st.dataframe(pd.DataFrame({"Column": df.columns, "Data Type": [str(t) for t in df.dtypes]}),
                 width="stretch", hide_index=True)
    st.subheader("First 100 Rows")
    st.dataframe(df.head(100), width="stretch", hide_index=True)
    st.download_button("Download Synthetic Scenario CSV", data=df.to_csv(index=False).encode("utf-8"),
                       file_name="macro_stress_synthetic_scenarios.csv", mime="text/csv",
                       key="download_synthetic_csv")


def tab_document(config):
    title("Document Analysis", "Institution classification and the configuration extracted from the document.")
    cur = config["currency"]
    cols = st.columns(4)
    amount = config.get("portfolio_amount")
    with cols[0]:
        kpi("File", st.session_state.upload_name or "N/A", "", "slate")
    with cols[1]:
        kpi("Institution", config["type"], "", "blue")
    with cols[2]:
        kpi("Portfolio", money(amount, cur) if amount else "Not detected", "", "green")
    with cols[3]:
        kpi("Scenario", config["scenario"]["name"], f"{config['scenario']['horizon_days']} days", "amber")
    st.markdown("")
    ok, errors = validate_config(config)
    if ok:
        st.success(f"{config['type']} institution detected. Configuration validation passed.")
    else:
        st.error("Configuration validation failed.")
        for e in errors:
            st.warning(e)
    st.subheader("Portfolio Assets")
    rows = []
    for name, a in config["portfolio_assets"].items():
        row = {"Asset": name, "Weight": pct(a["weight"]) if a.get("weight") is not None else "Missing"}
        for v in RISK_VARIABLES:
            row[f"{v} Sens."] = f"{a.get('sensitivities', {}).get(v, 0.0):+.4f}"
        rows.append(row)
    if rows:
        st.dataframe(pd.DataFrame(rows), width="stretch", hide_index=True)
    c1, c2 = st.columns(2)
    with c1:
        st.subheader("Constraints")
        st.dataframe(pd.DataFrame({"Constraint": ["Minimum weight", "Maximum weight"],
                                   "Value": [pct(config["constraints"]["min_weight"]),
                                             pct(config["constraints"]["max_weight"])]}),
                     width="stretch", hide_index=True)
    with c2:
        st.subheader("Stress Shocks")
        st.dataframe(pd.DataFrame({"Risk Variable": RISK_VARIABLES,
                                   "Shock": [shock_display(v, config["scenario"]["shocks"].get(v))
                                             for v in RISK_VARIABLES]}), width="stretch", hide_index=True)
    with st.expander("Extracted document text"):
        st.text_area("Document text", config.get("raw_text", ""), height=400, key="document_text_preview",
                     label_visibility="collapsed")


def tab_report(config, sim, r, bm, om, conf, checks):
    title("Final Report", "Export the completed MacroStress-GAN risk assessment.")
    html_report = build_report_html(config, sim, r, bm, om, conf, checks)
    c1, c2 = st.columns(2)
    c1.download_button("Download HTML Report", data=html_report.encode("utf-8"),
                       file_name="macro_stress_risk_report.html", mime="text/html", key="download_html_report")
    c2.download_button("Download JSON Report", data=export_json(config, r, bm, om, conf, checks),
                       file_name="macro_stress_risk_report.json", mime="application/json",
                       key="download_json_report")
    st.markdown("")
    st.subheader("Report Preview")
    st.components.v1.html(html_report, height=800, scrolling=True)


# ============================================================
# MAIN
# ============================================================
def main() -> None:
    run_clicked, conf = render_sidebar()
    config = st.session_state.config
    if config is None:
        landing()
        return

    hero(config["type"])
    valid, errors = validate_config(config)
    if not valid:
        st.error("The uploaded configuration is incomplete or invalid. Fix the document and upload it again.")
        tab_document(config)
        return

    if run_clicked:
        try:
            with st.status("Running dynamic pipeline...", expanded=True) as status:
                st.write("Generating TimeGAN paths and injecting the uploaded scenario...")
                sim = run_simulation(config)
                st.write("Optimizing the uploaded portfolio...")
                opt = run_optimization(config, sim)
                st.write("Re-stressing on the same paths...")
                rs = calculate_restress(config, sim, opt)
                status.update(label="Pipeline completed", state="complete")
            st.session_state.simulation, st.session_state.optimization, st.session_state.restress = sim, opt, rs
            st.session_state.pipeline_error = None
        except Exception as exc:
            st.session_state.pipeline_error = f"{type(exc).__name__}: {exc}"
            st.error("Pipeline execution failed.")
            st.code(traceback.format_exc())
            return

    sim, opt, rs = st.session_state.simulation, st.session_state.optimization, st.session_state.restress
    if sim is None or opt is None or rs is None:
        st.info("Configuration loaded and validated. Click **Run Dynamic Stress Test** in the sidebar to generate "
                "scenarios, risk metrics, the optimized portfolio and the re-stress verification.")
        tab_document(config)
        return

    value = float(config["portfolio_amount"])
    bm = risk_metrics(rs["base_daily"], value, conf)
    om = risk_metrics(rs["opt_daily"], value, conf)
    checks = restress_checks(rs, bm, om, config)
    raw = sim["stressed_raw"]
    rating = risk_rating(bm["Expected Loss"])

    c = st.columns(4)
    with c[0]:
        kpi("Institution", config["type"], "Detected from document", "blue")
    with c[1]:
        kpi("Scenarios", f"{raw.shape[0]:,}", "Generated stress paths", "teal")
    with c[2]:
        kpi("Horizon", f"{raw.shape[1]} days", "Per scenario", "blue")
    with c[3]:
        kpi("Risk Rating", rating, "Based on expected loss", RATING_ACCENT[rating])
    st.markdown("")

    tabs = st.tabs(["Overview", "Risk Analytics", "Optimization", "Re-Stress", "Data", "Document", "Report"])
    with tabs[0]:
        tab_overview(config, sim, rs, bm, conf)
    with tabs[1]:
        tab_risk(config, rs, bm, conf)
    with tabs[2]:
        tab_optimization(config, opt, rs, bm, om, conf)
    with tabs[3]:
        tab_restress(config, rs, bm, om, conf, checks)
    with tabs[4]:
        tab_data(sim)
    with tabs[5]:
        tab_document(config)
    with tabs[6]:
        tab_report(config, sim, rs, bm, om, conf, checks)


if __name__ == "__main__":
    main()