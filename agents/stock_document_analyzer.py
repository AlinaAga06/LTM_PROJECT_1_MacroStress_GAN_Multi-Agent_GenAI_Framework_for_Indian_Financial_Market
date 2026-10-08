"""
MacroStress-GAN
Stock Document Analyzer
============================================================

Purpose
-------
Safely analyze an institution portfolio configuration supplied
as a DOCX or PDF document.

The document is treated ONLY as DATA.
No Python/code from the document is executed.

Supported:
    - DOCX
    - PDF

Extracts:
    - Institution type
    - Portfolio amount
    - Currency
    - Risk variables
    - Portfolio assets
    - Asset weights
    - Asset sensitivities
    - Portfolio constraints
    - Stress scenario
    - Scenario shocks

Output:
    outputs/stock_document_analysis.json
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional


# ============================================================
# PATH CONFIGURATION
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]

DEFAULT_OUTPUT_PATH = (
    PROJECT_ROOT
    / "outputs"
    / "stock_document_analysis.json"
)


# ============================================================
# CONFIGURATION
# ============================================================

TIMEGAN_FEATURE_MAP = {
    "NIFTY50": "NIFTY50_Return",
    "CRUDE_OIL": "CRUDE_OIL_Return",
    "USD_INR": "USD_INR_Return",
    "INDIA_VIX": "INDIA_VIX_Change",
    "INDIA_10Y_YIELD": "INDIA_10Y_YIELD_Change",
}

SUPPORTED_RISK_VARIABLES = [
    "NIFTY50",
    "CRUDE_OIL",
    "USD_INR",
    "INDIA_VIX",
    "INDIA_10Y_YIELD",
]


# ============================================================
# BASIC TEXT HELPERS
# ============================================================

def normalize_text(value: Any) -> str:
    if value is None:
        return ""

    text = str(value)

    text = text.replace("\xa0", " ")
    text = text.replace("\u2013", "-")
    text = text.replace("\u2014", "-")
    text = text.replace("\u2212", "-")

    text = re.sub(r"[ \t]+", " ", text)

    return text.strip()


def normalize_label(value: Any) -> str:
    text = normalize_text(value).lower()
    text = re.sub(r"[^a-z0-9]+", "", text)
    return text


def clean_numeric_text(value: str) -> str:
    text = normalize_text(value)

    text = text.replace(",", "")
    text = text.replace("₹", "")
    text = text.replace("$", "")
    text = text.replace("€", "")
    text = text.replace("£", "")

    return text.strip()


def parse_number(value: Any) -> Optional[float]:
    if value is None:
        return None

    text = clean_numeric_text(str(value))

    match = re.search(
        r"[-+]?\d+(?:\.\d+)?",
        text,
    )

    if not match:
        return None

    try:
        return float(match.group(0))
    except ValueError:
        return None


def parse_percentage(value: Any) -> Optional[float]:
    if value is None:
        return None

    text = clean_numeric_text(str(value))

    number = parse_number(text)

    if number is None:
        return None

    if "%" in text:
        return number / 100.0

    if abs(number) <= 1.0:
        return number

    return number / 100.0


# ============================================================
# DOCUMENT EXTRACTION
# ============================================================

def extract_docx_lines(document_path: Path) -> List[str]:

    from docx import Document

    document = Document(str(document_path))

    lines: List[str] = []

    # Paragraphs
    for paragraph in document.paragraphs:

        text = normalize_text(paragraph.text)

        if text:
            lines.append(text)

    # Tables
    for table in document.tables:

        for row in table.rows:

            row_values = []

            for cell in row.cells:

                cell_text = normalize_text(cell.text)

                if cell_text:
                    row_values.append(cell_text)

            if not row_values:
                continue

            row_text = " | ".join(row_values)

            lines.append(row_text)

            for value in row_values:

                if value not in lines:
                    lines.append(value)

    return lines


def extract_pdf_lines(document_path: Path) -> List[str]:

    from pypdf import PdfReader

    reader = PdfReader(str(document_path))

    lines: List[str] = []

    for page in reader.pages:

        text = page.extract_text() or ""

        for raw_line in text.splitlines():

            line = normalize_text(raw_line)

            if line:
                lines.append(line)

    return lines


def extract_document_lines(
    document_path: Path,
) -> List[str]:

    suffix = document_path.suffix.lower()

    if suffix == ".docx":
        return extract_docx_lines(document_path)

    if suffix == ".pdf":
        return extract_pdf_lines(document_path)

    raise ValueError(
        f"Unsupported document format: {suffix}. "
        "Only .docx and .pdf are supported."
    )


# ============================================================
# LINE HELPERS
# ============================================================

def line_contains_label(
    line: str,
    label: str,
) -> bool:

    return normalize_label(label) in normalize_label(line)


def value_after_label(
    line: str,
    label: str,
) -> Optional[str]:

    text = normalize_text(line)

    pattern = re.compile(
        rf"^\s*{re.escape(label)}\s*[:\-|]\s*(.+?)\s*$",
        re.IGNORECASE,
    )

    match = pattern.match(text)

    if match:
        return match.group(1).strip()

    pattern2 = re.compile(
        rf"{re.escape(label)}\s*[:\-|]\s*(.+)",
        re.IGNORECASE,
    )

    match2 = pattern2.search(text)

    if match2:
        return match2.group(1).strip()

    return None


def find_labeled_value(
    lines: List[str],
    labels: List[str],
) -> Optional[str]:

    normalized_labels = [
        normalize_label(label)
        for label in labels
    ]

    for index, line in enumerate(lines):

        normalized_line = normalize_label(line)

        for label, normalized_label in zip(
            labels,
            normalized_labels,
        ):

            value = value_after_label(
                line,
                label,
            )

            if value:
                return value

            if normalized_line.startswith(
                normalized_label
            ):

                remainder = line[
                    len(label):
                ].strip(" :-|")

                if remainder:
                    return remainder

        for normalized_label in normalized_labels:

            if normalized_line == normalized_label:

                if index + 1 < len(lines):

                    next_line = normalize_text(
                        lines[index + 1]
                    )

                    if next_line:
                        return next_line

    return None


# ============================================================
# INSTITUTION
# ============================================================

def find_institution_type(
    lines: List[str],
) -> Optional[str]:

    return find_labeled_value(
        lines,
        [
            "Institution Type",
            "Institution",
            "Organization Type",
        ],
    )


def find_currency(
    lines: List[str],
) -> Optional[str]:

    value = find_labeled_value(
        lines,
        [
            "Currency",
            "Portfolio Currency",
        ],
    )

    if value:
        return value.strip().upper()

    return None


def find_portfolio_amount(
    lines: List[str],
) -> Optional[float]:

    value = find_labeled_value(
        lines,
        [
            "Portfolio Amount",
            "Portfolio Size",
            "Investment Amount",
            "Total Portfolio Amount",
            "Total Portfolio Size",
        ],
    )

    if value is None:
        return None

    return parse_number(value)


# ============================================================
# RISK VARIABLES
# ============================================================

def find_risk_variables(
    lines: List[str],
) -> List[str]:

    variables: List[str] = []

    for line in lines:

        normalized = normalize_label(line)

        for variable in SUPPORTED_RISK_VARIABLES:

            if normalize_label(variable) == normalized:

                if variable not in variables:
                    variables.append(variable)

    complete_text = " ".join(lines)

    for variable in SUPPORTED_RISK_VARIABLES:

        if normalize_label(variable) in normalize_label(
            complete_text
        ):

            if variable not in variables:
                variables.append(variable)

    return variables


# ============================================================
# PORTFOLIO ASSET DETECTION
# ============================================================

ASSET_PATTERN = re.compile(
    r"^\s*(Asset[_\s-]*[A-Za-z0-9]+)\s*$",
    re.IGNORECASE,
)


def canonical_asset_name(
    value: str,
) -> Optional[str]:

    text = normalize_text(value)

    match = ASSET_PATTERN.match(text)

    if not match:
        return None

    raw_name = match.group(1)

    suffix_match = re.search(
        r"Asset[_\s-]*([A-Za-z0-9]+)",
        raw_name,
        re.IGNORECASE,
    )

    if not suffix_match:
        return None

    suffix = suffix_match.group(1)

    return f"Asset_{suffix}"


def find_asset_positions(
    lines: List[str],
) -> List[tuple[str, int]]:

    positions: List[tuple[str, int]] = []

    for index, line in enumerate(lines):

        asset_name = canonical_asset_name(line)

        if asset_name:

            positions.append(
                (
                    asset_name,
                    index,
                )
            )

    result: List[tuple[str, int]] = []

    seen = set()

    for asset_name, index in positions:

        key = (asset_name, index)

        if key not in seen:

            seen.add(key)

            result.append(
                (
                    asset_name,
                    index,
                )
            )

    return result


# ============================================================
# ASSET FIELD PARSING
# ============================================================

def extract_weight_from_block(
    block: List[str],
) -> Optional[float]:

    for line in block:

        value = value_after_label(
            line,
            "Weight",
        )

        if value is not None:
            return parse_percentage(value)

        match = re.search(
            r"\bWeight\b\s*[:\-]?\s*"
            r"([-+]?\d+(?:\.\d+)?)\s*%",
            line,
            re.IGNORECASE,
        )

        if match:

            return (
                float(match.group(1))
                / 100.0
            )

    return None


def extract_sensitivity(
    line: str,
    variable: str,
) -> Optional[float]:

    variable_pattern = re.escape(variable)

    pattern = re.compile(
        rf"{variable_pattern}\s+Sensitivity\s*[:\-|]\s*"
        rf"([-+]?\d+(?:\.\d+)?)",
        re.IGNORECASE,
    )

    match = pattern.search(line)

    if match:
        return float(match.group(1))

    normalized_line = normalize_label(line)
    normalized_variable = normalize_label(variable)

    if (
        normalized_variable in normalized_line
        and "sensitivity" in normalized_line
    ):

        number = parse_number(line)

        if number is not None:
            return number

    return None


def extract_asset_sensitivities(
    block: List[str],
) -> Dict[str, float]:

    sensitivities: Dict[str, float] = {}

    for line in block:

        for variable in SUPPORTED_RISK_VARIABLES:

            value = extract_sensitivity(
                line,
                variable,
            )

            if value is not None:

                sensitivities[variable] = value

    return sensitivities


def parse_portfolio_assets(
    lines: List[str],
) -> Dict[str, Dict[str, Any]]:

    positions = find_asset_positions(lines)

    assets: Dict[str, Dict[str, Any]] = {}

    for position_index, (
        asset_name,
        start_index,
    ) in enumerate(positions):

        if position_index + 1 < len(positions):

            end_index = positions[
                position_index + 1
            ][1]

        else:

            end_index = len(lines)

        block = lines[
            start_index:end_index
        ]

        weight = extract_weight_from_block(
            block
        )

        sensitivities = extract_asset_sensitivities(
            block
        )

        assets[asset_name] = {
            "weight": (
                weight
                if weight is not None
                else 0.0
            ),
            "sensitivities": sensitivities,
        }

    return assets


# ============================================================
# CONSTRAINTS
# ============================================================

def find_minimum_weight(
    lines: List[str],
) -> Optional[float]:

    value = find_labeled_value(
        lines,
        [
            "Minimum Weight",
            "Min Weight",
            "Minimum Portfolio Weight",
        ],
    )

    if value is None:
        return None

    return parse_percentage(value)


def find_maximum_weight(
    lines: List[str],
) -> Optional[float]:

    value = find_labeled_value(
        lines,
        [
            "Maximum Weight",
            "Max Weight",
            "Maximum Portfolio Weight",
        ],
    )

    if value is None:
        return None

    return parse_percentage(value)


# ============================================================
# SCENARIO
# ============================================================

def find_scenario_name(
    lines: List[str],
) -> Optional[str]:

    value = find_labeled_value(
        lines,
        [
            "Scenario Name",
            "Stress Scenario Name",
            "Scenario",
        ],
    )

    if not value:
        return None

    return value.strip(" |:-").strip()


def find_horizon_days(
    lines: List[str],
) -> int:

    value = find_labeled_value(
        lines,
        [
            "Horizon",
            "Horizon Days",
            "Stress Horizon",
        ],
    )

    if value is None:
        return 30

    match = re.search(
        r"(\d+)",
        value,
    )

    if match:
        return int(match.group(1))

    return 30


def find_shock(
    lines: List[str],
    labels: List[str],
) -> float:

    value = find_labeled_value(
        lines,
        labels,
    )

    if value is None:
        return 0.0

    parsed = parse_percentage(value)

    if parsed is None:
        return 0.0

    return parsed


def find_scenario_shocks(
    lines: List[str],
) -> Dict[str, float]:

    return {

        "NIFTY50": find_shock(
            lines,
            [
                "NIFTY50 Shock",
                "NIFTY50 Market Shock",
            ],
        ),

        "CRUDE_OIL": find_shock(
            lines,
            [
                "CRUDE_OIL Shock",
                "CRUDE OIL Shock",
                "Crude Oil Shock",
            ],
        ),

        "USD_INR": find_shock(
            lines,
            [
                "USD_INR Shock",
                "USD/INR Shock",
                "USD INR Shock",
            ],
        ),

        "INDIA_VIX": find_shock(
            lines,
            [
                "INDIA_VIX Shock",
                "India VIX Shock",
                "VIX Shock",
            ],
        ),

        "INDIA_10Y_YIELD": find_shock(
            lines,
            [
                "INDIA_10Y_YIELD Shock",
                "India 10Y Yield Shock",
                "10Y Yield Shock",
            ],
        ),
    }


# ============================================================
# VALIDATION
# ============================================================

def validate_weights(
    assets: Dict[str, Dict[str, Any]],
) -> Dict[str, Any]:

    if not assets:

        return {
            "weight_total": 0.0,
            "weights_valid": False,
            "all_weights_present": False,
        }

    weight_total = sum(
        float(
            asset.get(
                "weight",
                0.0,
            )
        )
        for asset in assets.values()
    )

    all_weights_present = all(
        asset.get("weight") is not None
        and float(
            asset.get(
                "weight",
                0.0,
            )
        ) > 0
        for asset in assets.values()
    )

    weights_valid = (
        all_weights_present
        and abs(weight_total - 1.0)
        <= 1e-6
    )

    return {
        "weight_total": weight_total,
        "weights_valid": weights_valid,
        "all_weights_present": all_weights_present,
    }


def validate_constraints(
    constraints: Dict[str, Any],
) -> bool:

    minimum = constraints.get(
        "min_weight"
    )

    maximum = constraints.get(
        "max_weight"
    )

    if minimum is None or maximum is None:
        return False

    if minimum < 0:
        return False

    if maximum > 1:
        return False

    if minimum > maximum:
        return False

    return True


def validate_scenario(
    scenario: Dict[str, Any],
) -> bool:

    if not scenario:
        return False

    if not scenario.get(
        "scenario_name"
    ):
        return False

    if not scenario.get(
        "horizon_days"
    ):
        return False

    shocks = scenario.get(
        "shocks",
        {},
    )

    required = set(
        SUPPORTED_RISK_VARIABLES
    )

    return required.issubset(
        set(shocks.keys())
    )


# ============================================================
# MAIN ANALYZER
# ============================================================

def analyze_document(
    document_path: str | Path,
    output_path: str | Path = DEFAULT_OUTPUT_PATH,
) -> Dict[str, Any]:

    document_path = Path(
        document_path
    ).resolve()

    output_path = Path(
        output_path
    ).resolve()

    if not document_path.exists():

        raise FileNotFoundError(
            f"Document not found: {document_path}"
        )

    lines = extract_document_lines(
        document_path
    )

    if not lines:

        raise ValueError(
            "No readable text was extracted "
            "from the document."
        )

    # --------------------------------------------------------
    # Institution
    # --------------------------------------------------------

    institution_type = find_institution_type(
        lines
    )

    currency = find_currency(
        lines
    )

    portfolio_amount = find_portfolio_amount(
        lines
    )

    # --------------------------------------------------------
    # Risk Variables
    # --------------------------------------------------------

    risk_variables = find_risk_variables(
        lines
    )

    if not risk_variables:

        risk_variables = list(
            SUPPORTED_RISK_VARIABLES
        )

    timegan_features = [
        TIMEGAN_FEATURE_MAP[var]
        for var in risk_variables
        if var in TIMEGAN_FEATURE_MAP
    ]

    # --------------------------------------------------------
    # Portfolio Assets
    # --------------------------------------------------------

    portfolio_assets = parse_portfolio_assets(
        lines
    )

    # --------------------------------------------------------
    # Constraints
    # --------------------------------------------------------

    min_weight = find_minimum_weight(
        lines
    )

    max_weight = find_maximum_weight(
        lines
    )

    if min_weight is None:
        min_weight = 0.05

    if max_weight is None:
        max_weight = 0.60

    constraints = {
        "min_weight": min_weight,
        "max_weight": max_weight,
    }

    # --------------------------------------------------------
    # Scenario
    # --------------------------------------------------------

    scenario_name = find_scenario_name(
        lines
    )

    if not scenario_name:
        scenario_name = (
            "Institutional Stress Test"
        )

    horizon_days = find_horizon_days(
        lines
    )

    shocks = find_scenario_shocks(
        lines
    )

    scenario = {
        "scenario_name": scenario_name,
        "horizon_days": horizon_days,
        "shocks": shocks,
    }

    # --------------------------------------------------------
    # Validation
    # --------------------------------------------------------

    weight_validation = validate_weights(
        portfolio_assets
    )

    constraints_valid = validate_constraints(
        constraints
    )

    scenario_valid = validate_scenario(
        scenario
    )

    errors: List[str] = []

    if portfolio_amount is None:

        errors.append(
            "Portfolio amount was not found."
        )

    if not portfolio_assets:

        errors.append(
            "No portfolio assets were found."
        )

    if not weight_validation[
        "all_weights_present"
    ]:

        errors.append(
            "No portfolio asset weights "
            "were found."
        )

    if (
        not weight_validation[
            "weights_valid"
        ]
        and portfolio_assets
    ):

        errors.append(
            "Portfolio asset weights "
            "must sum to 100%."
        )

    if not constraints_valid:

        errors.append(
            "Portfolio constraints are invalid."
        )

    if not scenario_valid:

        errors.append(
            "Stress scenario is invalid."
        )

    status = (
        "validated"
        if not errors
        else "error"
    )

    # --------------------------------------------------------
    # Result
    # --------------------------------------------------------

    result: Dict[str, Any] = {

        "status": status,

        "document": {
            "path": str(document_path),
            "format": document_path.suffix.lower(),
        },

        "institution": {

            "type":
                institution_type
                or "Stock Market",

            "portfolio_amount":
                portfolio_amount,

            "currency":
                currency
                or "INR",
        },

        "institution_type":
            institution_type
            or "Stock Market",

        "portfolio_amount":
            portfolio_amount,

        "currency":
            currency
            or "INR",

        "risk_variables":
            risk_variables,

        "timegan_features":
            timegan_features,

        "portfolio_assets":
            portfolio_assets,

        "constraints":
            constraints,

        "scenario":
            scenario,

        "validation": {

            "weight_total":
                weight_validation[
                    "weight_total"
                ],

            "weight_total_percent":
                weight_validation[
                    "weight_total"
                ] * 100.0,

            "weights_valid":
                weight_validation[
                    "weights_valid"
                ],

            "constraints_valid":
                constraints_valid,

            "scenario_valid":
                scenario_valid,
        },

        "errors":
            errors,
    }

    # --------------------------------------------------------
    # Save JSON
    # --------------------------------------------------------

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with open(
        output_path,
        "w",
        encoding="utf-8",
    ) as file:

        json.dump(
            result,
            file,
            indent=2,
            ensure_ascii=False,
        )

    # --------------------------------------------------------
    # Console Output
    # --------------------------------------------------------

    print()
    print("=" * 80)
    print("MACROSTRESS-GAN")
    print("STOCK DOCUMENT ANALYZER")
    print("=" * 80)

    print(
        f"STATUS: {status}"
    )

    print()
    print("DOCUMENT")

    print(
        f"Path: {document_path}"
    )

    print(
        f"Format: "
        f"{document_path.suffix.lower()}"
    )

    print()
    print("INSTITUTION")

    print(
        f"Type: "
        f"{institution_type or 'Stock Market'}"
    )

    if portfolio_amount is None:

        print(
            "Portfolio Amount: None"
        )

    else:

        print(
            f"Portfolio Amount: "
            f"{portfolio_amount:,.2f}"
        )

    print(
        f"Currency: "
        f"{currency or 'INR'}"
    )

    print()
    print("RISK VARIABLES")

    for variable in risk_variables:

        print(
            f"  {variable}"
        )

    print()
    print("TIMEGAN FEATURES")

    for feature in timegan_features:

        print(
            f"  {feature}"
        )

    print()
    print("PORTFOLIO ASSETS")

    if portfolio_assets:

        for asset_name, asset in (
            portfolio_assets.items()
        ):

            weight = float(
                asset.get(
                    "weight",
                    0.0,
                )
            )

            print(
                f"  {asset_name}: "
                f"{weight:.4f} "
                f"({weight * 100:.2f}%)"
            )

            sensitivities = asset.get(
                "sensitivities",
                {},
            )

            for variable, sensitivity in (
                sensitivities.items()
            ):

                print(
                    f"    {variable}: "
                    f"{sensitivity:+.4f}"
                )

    else:

        print(
            "  No portfolio assets found."
        )

    print()
    print("CONSTRAINTS")

    print(
        f"Minimum Weight: "
        f"{min_weight * 100:.2f}%"
    )

    print(
        f"Maximum Weight: "
        f"{max_weight * 100:.2f}%"
    )

    print()
    print("STRESS SCENARIO")

    print(
        f"  Scenario Name: "
        f"{scenario_name}"
    )

    print(
        f"  Horizon: "
        f"{horizon_days} days"
    )

    print(
        "  Market Shocks:"
    )

    for variable in SUPPORTED_RISK_VARIABLES:

        shock = shocks.get(
            variable,
            0.0,
        )

        print(
            f"    {variable}: "
            f"{shock * 100:+.2f}%"
        )

    print()
    print("VALIDATION")

    print(
        f"Weight Total: "
        f"{weight_validation['weight_total'] * 100:.4f}%"
    )

    print(
        f"Weights Valid: "
        f"{weight_validation['weights_valid']}"
    )

    print(
        f"Constraints Valid: "
        f"{constraints_valid}"
    )

    print(
        f"Scenario Valid: "
        f"{scenario_valid}"
    )

    if errors:

        print()
        print("ERRORS")

        for error in errors:

            print(
                f"  - {error}"
            )

    print()
    print(
        f"Output: {output_path}"
    )

    print("=" * 80)

    return result


# ============================================================
# CLI
# ============================================================

def main() -> int:

    if len(sys.argv) < 2:

        print("Usage:")

        print(
            "python .\\agents\\document_analysis_agent.py "
            ".\\institution_portfolio.docx"
        )

        return 1

    document_path = Path(
        sys.argv[1]
    )

    try:

        result = analyze_document(
            document_path=document_path
        )

        return (
            0
            if result["status"]
            == "validated"
            else 1
        )

    except Exception as exc:

        print()
        print("STATUS: error")

        print(
            f"ERROR: {exc}"
        )

        return 1


if __name__ == "__main__":

    raise SystemExit(
        main()
    )